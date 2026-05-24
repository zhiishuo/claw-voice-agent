#!/usr/bin/env python3
import csv
import hashlib
import json
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".jsonl",
    ".csv",
    ".xlsx",
    ".pdf",
    ".docx",
}


@dataclass
class Document:
    text: str
    metadata: dict = field(default_factory=dict)


def iter_source_files(source_dir: Path) -> Iterable[Path]:
    if source_dir.is_file():
        if source_dir.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield source_dir
        return
    for path in sorted(source_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield path


def load_documents(source_dir: str | Path) -> list[Document]:
    root = Path(source_dir).expanduser()
    if not root.exists():
        raise FileNotFoundError(f"knowledge source path does not exist: {root}")

    documents: list[Document] = []
    seen_checksums: set[str] = set()
    for path in iter_source_files(root):
        loaded = load_document(path, root if root.is_dir() else path.parent)
        if not loaded:
            continue
        checksum = loaded[0].metadata.get("file_checksum")
        if checksum in seen_checksums:
            continue
        seen_checksums.add(checksum)
        documents.extend(loaded)
    return documents


def load_document(path: Path, root: Path) -> list[Document]:
    suffix = path.suffix.lower()
    metadata = _file_metadata(path, root)
    if suffix in {".txt", ".md", ".markdown"}:
        return [Document(path.read_text(encoding="utf-8", errors="ignore"), metadata)]
    if suffix == ".json":
        return [Document(_read_json(path), metadata)]
    if suffix == ".jsonl":
        return _read_jsonl(path, metadata)
    if suffix == ".csv":
        return [Document(_read_csv(path), {**metadata, "chunk_type": "table"})]
    if suffix == ".xlsx":
        return _read_xlsx(path, metadata)
    if suffix == ".pdf":
        return [Document(_read_pdf(path), {**metadata, "chunk_type": "document"})]
    if suffix == ".docx":
        return [Document(_read_docx(path), {**metadata, "chunk_type": "document"})]
    return []


def _file_metadata(path: Path, root: Path) -> dict:
    stat = path.stat()
    checksum = _sha256(path)
    relative_path = str(path.relative_to(root)) if path.is_relative_to(root) else path.name
    return {
        "source": str(path),
        "filename": path.name,
        "extension": path.suffix.lower(),
        "relative_path": relative_path,
        "file_checksum": checksum,
        "file_size": stat.st_size,
        "file_mtime": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
        "ingest_time": datetime.now(timezone.utc).isoformat(),
        "document_id": checksum[:16],
    }


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return json.dumps(data, ensure_ascii=False, indent=2)


def _read_jsonl(path: Path, metadata: dict) -> list[Document]:
    docs: list[Document] = []
    with path.open("r", encoding="utf-8", errors="ignore") as fh:
        for line_no, line in enumerate(fh, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                payload = json.loads(line)
                text = json.dumps(payload, ensure_ascii=False)
            except json.JSONDecodeError:
                text = line
            docs.append(Document(text, {**metadata, "line": line_no, "chunk_type": "record"}))
    return docs


def _read_csv(path: Path) -> str:
    rows: list[str] = []
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames:
            rows.append("fields: " + " | ".join(reader.fieldnames))
            for row_number, row in enumerate(reader, start=2):
                cells = [f"{key}: {value}" for key, value in row.items() if value]
                if cells:
                    rows.append(f"row {row_number}: " + " | ".join(cells))
        else:
            fh.seek(0)
            for row_number, row in enumerate(csv.reader(fh), start=1):
                cells = [cell for cell in row if cell]
                if cells:
                    rows.append(f"row {row_number}: " + " | ".join(cells))
    return "\n".join(rows)


def _read_xlsx(path: Path, metadata: dict) -> list[Document]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX parsing requires openpyxl: pip install openpyxl") from exc

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    documents: list[Document] = []
    try:
        for sheet in workbook.worksheets:
            header_row_number, header, leading_rows = _find_xlsx_header(sheet.iter_rows(values_only=True))
            if header is None:
                continue

            headers = [
                _cell_to_text(value) or f"column_{index}"
                for index, value in enumerate(header, start=1)
            ]
            lines = [f"[sheet: {sheet.title}]", "fields: " + " | ".join(headers)]
            for row_number, values in leading_rows:
                lines.append(f"row {row_number}: " + " | ".join(values))

            for row_number, row in enumerate(
                sheet.iter_rows(min_row=header_row_number + 1, values_only=True),
                start=header_row_number + 1,
            ):
                values = [_cell_to_text(value) for value in row]
                if not any(values):
                    continue
                cells = [
                    f"{headers[index]}: {value}"
                    for index, value in enumerate(values)
                    if value and index < len(headers)
                ]
                if cells:
                    lines.append(f"row {row_number}: " + " | ".join(cells))
            documents.append(
                Document(
                    "\n".join(lines),
                    {
                        **metadata,
                        "sheet_name": sheet.title,
                        "table_fields": headers,
                        "chunk_type": "table",
                    },
                )
            )
    finally:
        workbook.close()

    return documents


def _find_xlsx_header(rows) -> tuple[int, tuple[str, ...] | None, list[tuple[int, list[str]]]]:
    leading_rows: list[tuple[int, list[str]]] = []
    fallback: tuple[int, tuple[str, ...]] | None = None
    for row_number, row in enumerate(rows, start=1):
        values = tuple(_cell_to_text(value) for value in row)
        non_empty = [value for value in values if value]
        if not non_empty:
            continue
        if fallback is None:
            fallback = (row_number, values)
        if len(non_empty) >= 2:
            return row_number, values, leading_rows
        leading_rows.append((row_number, non_empty))
    if fallback is None:
        return 0, None, []
    return fallback[0], fallback[1], []


def _cell_to_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF parsing requires pypdf: pip install pypdf") from exc

    reader = PdfReader(str(path))
    pages = []
    for page_no, page in enumerate(reader.pages, start=1):
        pages.append(f"[page {page_no}]\n{page.extract_text() or ''}")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    if _looks_like_legacy_doc(path):
        raise RuntimeError(
            f"{path} has a legacy Office binary signature. Convert it to a real .docx before indexing."
        )
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("DOCX parsing requires python-docx: pip install python-docx") from exc

    document = docx.Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.lower().startswith("heading"):
            lines.append(f"# {text}")
        else:
            lines.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        lines.append(f"[table {table_index}]")
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(f"row {row_index}: " + " | ".join(cells))

    return "\n".join(lines)


def _looks_like_legacy_doc(path: Path) -> bool:
    with path.open("rb") as fh:
        return fh.read(4) == b"\xd0\xcf\x11\xe0"
